from __future__ import annotations

import io
import json
from collections import Counter
from datetime import datetime
from typing import Mapping

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from schemas.analysis import AnalysisResult
from schemas.lesson_plan import AnalysisPackage, LessonPlanProfile
from schemas.package import (
    AnswerKeyProfile,
    ComparisonStatus,
    PackageComparisonResult,
    StudentQuestionSet,
    TeachingContextProfile,
)
from schemas.usage import AIUsageSummary
from services.package_result_service import (
    ACTIONABLE_STATUSES,
    package_decision,
    package_statistics,
)


_COMPARISON_STATUS_LABELS = {
    ComparisonStatus.CONFLICT: "Cần xác minh xung đột",
    ComparisonStatus.AMBIGUOUS: "Chưa rõ",
    ComparisonStatus.MISSING_KEY: "Thiếu đáp án đã xác nhận",
    ComparisonStatus.KEY_EXTRACTION_UNCERTAIN: "Chưa đọc được đáp án",
    ComparisonStatus.VERSION_MISMATCH: "Khác phiên bản",
}


def create_json_report(
    result: AnalysisResult,
    *,
    lesson_plan: LessonPlanProfile | None = None,
    lesson_plan_file_name: str | None = None,
    learning_material_file_name: str | None = None,
) -> bytes:
    payload = (
        AnalysisPackage(
            lesson_plan_file_name=lesson_plan_file_name or "lesson-plan",
            learning_material_file_name=learning_material_file_name or "learning-material",
            lesson_plan=lesson_plan,
            learning_material_analysis=result,
        ).model_dump(mode="json")
        if lesson_plan is not None
        else result.model_dump(mode="json")
    )
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def _add_bullets(document: Document, title: str, values: list[str]) -> None:
    if not values:
        return
    if title:
        document.add_paragraph(title)
    for value in values:
        document.add_paragraph(value, style="List Bullet")


def create_docx_report(
    result: AnalysisResult,
    file_name: str,
    metadata: Mapping[str, str],
    analyzed_at: datetime | None = None,
    *,
    lesson_plan: LessonPlanProfile | None = None,
    lesson_plan_file_name: str | None = None,
) -> bytes:
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    title = document.add_heading("BÁO CÁO KIỂM TRA HỌC LIỆU", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    analyzed_at = analyzed_at or datetime.now().astimezone()
    document.add_paragraph(f"Học liệu: {file_name}")
    if lesson_plan_file_name:
        document.add_paragraph(f"Giáo án đối chiếu: {lesson_plan_file_name}")
    document.add_paragraph(f"Thời điểm phân tích: {analyzed_at:%d/%m/%Y %H:%M}")
    if metadata.get("custom_request"):
        document.add_paragraph(f"Yêu cầu riêng: {metadata['custom_request']}")

    section = 1
    if lesson_plan is not None:
        document.add_heading(f"{section}. Căn cứ từ giáo án", level=1)
        section += 1
        document.add_paragraph(lesson_plan.summary)
        for label, value in (
            ("Môn học", lesson_plan.subject),
            ("Khối/lớp", lesson_plan.grade_level),
            ("Chủ đề", lesson_plan.topic),
        ):
            if value:
                document.add_paragraph(f"{label}: {value}")
        _add_bullets(document, "Mục tiêu bài học:", lesson_plan.objectives)
        _add_bullets(document, "Phạm vi kiến thức:", lesson_plan.knowledge_scope)
        _add_bullets(document, "Tiêu chí đánh giá:", lesson_plan.assessment_criteria)
        _add_bullets(
            document,
            "Thông tin giáo án cần xác minh:",
            lesson_plan.uncertainties,
        )

    summary = result.document_summary
    document.add_heading(f"{section}. Tóm tắt học liệu", level=1)
    section += 1
    document.add_paragraph(summary.summary)
    document.add_paragraph(
        f"Loại tài liệu: {summary.document_type} | "
        f"Ngôn ngữ: {', '.join(summary.detected_languages) or 'Không xác định'}"
    )
    if summary.total_questions is not None:
        document.add_paragraph(f"Tổng số câu hỏi: {summary.total_questions}")

    document.add_heading(f"{section}. Thống kê vấn đề", level=1)
    section += 1
    severity_counts = Counter(issue.severity.value for issue in result.issues)
    document.add_paragraph(
        f"Tổng cộng: {len(result.issues)} | "
        f"Nghiêm trọng: {severity_counts['critical']} | "
        f"Cảnh báo: {severity_counts['warning']} | "
        f"Gợi ý: {severity_counts['suggestion']}"
    )

    document.add_heading(f"{section}. Danh sách vấn đề trong học liệu", level=1)
    section += 1
    if not result.issues:
        document.add_paragraph("Không phát hiện vấn đề trong phạm vi đã chọn.")
    for index, issue in enumerate(result.issues, start=1):
        location = (
            f"Trang {issue.page_number}"
            if issue.page_number
            else issue.block_id or issue.section_title or "Không xác định"
        )
        document.add_heading(
            f"{index}. {issue.category.value} — {issue.severity.value}",
            level=2,
        )
        document.add_paragraph(
            f"Vị trí: {location}"
            + (f" | Câu: {issue.question_number}" if issue.question_number else "")
            + f" | Độ tin cậy: {issue.confidence:.0%}"
        )
        document.add_paragraph(f"Nội dung gốc: {issue.original_text}")
        document.add_paragraph(f"Giải thích: {issue.explanation}")
        if issue.lesson_plan_reference:
            document.add_paragraph(
                f"Căn cứ từ giáo án: {issue.lesson_plan_reference}"
            )
        if issue.suggested_text:
            document.add_paragraph(f"Đề xuất sửa: {issue.suggested_text}")
        if issue.replacement_question:
            document.add_paragraph(f"Câu hỏi thay thế: {issue.replacement_question}")
        if issue.suggested_answer:
            document.add_paragraph(f"Đáp án gợi ý: {issue.suggested_answer}")
        if issue.requires_human_review:
            document.add_paragraph("Lưu ý: Cần người dùng xác minh.")

    document.add_heading(f"{section}. Đề xuất chung", level=1)
    section += 1
    _add_bullets(document, "", result.general_recommendations)
    if not result.general_recommendations:
        document.add_paragraph("Không có đề xuất chung.")

    document.add_heading(f"{section}. Giới hạn", level=1)
    _add_bullets(document, "", result.limitations)
    document.add_paragraph(
        "Kết quả AI chỉ mang tính hỗ trợ và không thay thế thẩm định của giáo viên."
    )

    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def create_package_json_report(
    *,
    result: AnalysisResult,
    lesson_plan: LessonPlanProfile,
    teaching_context: TeachingContextProfile,
    questions: StudentQuestionSet,
    answer_key: AnswerKeyProfile,
    comparison: PackageComparisonResult,
    lesson_plan_file_name: str,
    teaching_file_names: list[str],
    learning_material_file_name: str,
    answer_key_file_name: str,
    metadata: Mapping[str, str] | None = None,
    usage: AIUsageSummary | None = None,
) -> bytes:
    """Create the complete, auditable result for a four-role learning package."""

    payload = {
        "package_files": {
            "lesson_plan": lesson_plan_file_name,
            "teaching_materials": teaching_file_names,
            "student_material": learning_material_file_name,
            "answer_key": answer_key_file_name,
        },
        "analysis_metadata": dict(metadata or {}),
        "ai_usage": usage.model_dump(mode="json") if usage else None,
        "lesson_plan": lesson_plan.model_dump(mode="json"),
        "teaching_context": teaching_context.model_dump(mode="json"),
        "student_questions": questions.model_dump(mode="json"),
        "answer_key": answer_key.model_dump(mode="json"),
        "comparison": comparison.model_dump(mode="json"),
        "student_material_analysis": result.model_dump(mode="json"),
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def _add_label_value(document: Document, label: str, value: str | None) -> None:
    if value:
        document.add_paragraph(f"{label}: {value}")


def create_package_docx_report(
    *,
    result: AnalysisResult,
    lesson_plan: LessonPlanProfile,
    teaching_context: TeachingContextProfile,
    questions: StudentQuestionSet,
    answer_key: AnswerKeyProfile,
    comparison: PackageComparisonResult,
    lesson_plan_file_name: str,
    teaching_file_names: list[str],
    learning_material_file_name: str,
    answer_key_file_name: str,
    metadata: Mapping[str, str],
    usage: AIUsageSummary | None = None,
    analyzed_at: datetime | None = None,
) -> bytes:
    """Create a teacher-friendly report for the complete learning package."""

    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    title = document.add_heading("BÁO CÁO KIỂM TRA BỘ HỌC LIỆU", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    analyzed_at = analyzed_at or datetime.now().astimezone()
    document.add_paragraph(f"Thời điểm phân tích: {analyzed_at:%d/%m/%Y %H:%M}")
    document.add_paragraph(f"Giáo án: {lesson_plan_file_name}")
    document.add_paragraph(
        "Tài liệu giảng dạy: " + ", ".join(teaching_file_names)
    )
    document.add_paragraph(f"Bài tập học viên: {learning_material_file_name}")
    document.add_paragraph(f"Đáp án: {answer_key_file_name}")
    if metadata.get("custom_request"):
        document.add_paragraph(f"Yêu cầu riêng: {metadata['custom_request']}")
    document.add_paragraph(
        "Chế độ: "
        + (
            "Phân tích toàn diện"
            if metadata.get("analysis_mode") == "full"
            else "Kiểm tra nhanh"
        )
    )
    if usage is not None:
        document.add_paragraph(
            f"Chi phí AI ước tính: ${usage.estimated_cost_usd:.4f} | "
            f"Lượt gọi API: {usage.call_count} | "
            f"Input token: {usage.input_tokens:,} | "
            f"Output token: {usage.output_tokens:,}"
        )
        document.add_paragraph(
            "Chi phí được ước tính theo giá Standard API và có thể khác hóa đơn thực tế."
        )

    document.add_heading("1. Căn cứ bài dạy", level=1)
    document.add_paragraph(lesson_plan.summary)
    _add_label_value(document, "Môn học", lesson_plan.subject)
    _add_label_value(document, "Khối/lớp", lesson_plan.grade_level)
    _add_label_value(document, "Chủ đề", lesson_plan.topic)
    _add_bullets(document, "Mục tiêu bài học:", lesson_plan.objectives)
    _add_bullets(document, "Phạm vi kiến thức:", lesson_plan.knowledge_scope)
    _add_bullets(document, "Tiêu chí đánh giá:", lesson_plan.assessment_criteria)

    document.add_heading("2. Nội dung thực tế đã dạy", level=1)
    document.add_paragraph(teaching_context.summary)
    _add_bullets(document, "Nội dung đã bao phủ:", teaching_context.covered_content)
    _add_bullets(document, "Thuật ngữ chính:", teaching_context.key_terms)
    _add_bullets(document, "Ví dụ đã hướng dẫn:", teaching_context.worked_examples)
    _add_bullets(
        document,
        "Quy ước và hướng dẫn:",
        teaching_context.instructions_and_conventions,
    )
    _add_bullets(
        document,
        "Điểm cần xác minh:",
        lesson_plan.uncertainties + teaching_context.uncertainties,
    )

    document.add_heading("3. Kết quả đối chiếu bài tập và đáp án", level=1)
    document.add_paragraph(comparison.summary)
    stats = package_statistics(questions, answer_key, comparison)
    decision, decision_detail = package_decision(stats)
    document.add_heading(f"Kết luận: {decision}", level=2)
    document.add_paragraph(decision_detail)
    document.add_paragraph(
        f"Câu trong bài tập: {stats.question_count} | "
        f"Khớp: {stats.match_count} | "
        f"AI trả lời khác: {stats.ai_disagreement_count} | "
        f"Cần xác minh: {stats.actionable_count} | "
        f"Ngoài phạm vi: {stats.outside_scope_count}"
    )
    actionable_items = [
        item for item in comparison.comparisons if item.status in ACTIONABLE_STATUSES
    ]
    if not actionable_items:
        document.add_paragraph("Không có mục đối chiếu nào cần giáo viên xác minh.")
    for index, item in enumerate(actionable_items, start=1):
        document.add_heading(
            f"{index}. Câu {item.question_number} — "
            f"{_COMPARISON_STATUS_LABELS.get(item.status, item.status.value)}",
            level=2,
        )
        document.add_paragraph(
            f"AI tự giải: {item.independent_answer or 'Không xác định'}"
        )
        document.add_paragraph(
            f"Đáp án chính thức: {item.official_answer or 'Không có'}"
        )
        document.add_paragraph(f"Nhận xét: {item.explanation}")
        _add_label_value(document, "Căn cứ giáo án", item.lesson_plan_reference)
        _add_label_value(
            document,
            "Căn cứ tài liệu dạy",
            item.teaching_material_reference,
        )
        document.add_paragraph(f"Độ tin cậy: {item.confidence:.0%}")
        if item.requires_human_review:
            document.add_paragraph("Lưu ý: Cần giáo viên xác minh.")
    _add_bullets(document, "Đề xuất xử lý:", comparison.recommendations)

    document.add_heading("4. Câu hỏi và đáp án đã trích xuất", level=1)
    document.add_paragraph(
        f"AI nhận diện {len(questions.questions)} câu hỏi và "
        f"{len(answer_key.entries)} mục đáp án."
    )
    for question in questions.questions:
        location = (
            f"trang {question.page_number}"
            if question.page_number
            else question.block_id or "chưa xác định"
        )
        document.add_heading(
            f"Câu {question.question_number} ({location})",
            level=2,
        )
        document.add_paragraph(question.question_text)
        _add_bullets(document, "", question.options)
        document.add_paragraph(
            f"AI tự giải: {question.independent_answer or 'Không xác định'}"
        )
        document.add_paragraph(f"Lập luận tóm tắt: {question.reasoning_summary}")
        if question.evidence_ids:
            document.add_paragraph(
                "Dẫn chứng đã dùng: " + ", ".join(question.evidence_ids)
            )
        for evidence in question.evidence:
            document.add_paragraph(
                f"{evidence.evidence_id} — {evidence.source}, "
                f"{evidence.location}: {evidence.excerpt}",
                style="List Bullet",
            )

    document.add_heading("5. Kiểm tra chi tiết bản thân bài tập", level=1)
    if metadata.get("analysis_mode") != "full":
        document.add_paragraph(
            "Chưa chạy kiểm tra chi tiết chất lượng tiếng Anh trong chế độ nhanh."
        )
    else:
        severity_counts = Counter(issue.severity.value for issue in result.issues)
        document.add_paragraph(
            f"Tổng cộng: {len(result.issues)} | "
            f"Nghiêm trọng: {severity_counts['critical']} | "
            f"Cảnh báo: {severity_counts['warning']} | "
            f"Gợi ý: {severity_counts['suggestion']}"
        )
        if not result.issues:
            document.add_paragraph("Không phát hiện vấn đề trong phạm vi đã chọn.")
        for index, issue in enumerate(result.issues, start=1):
            location = (
                f"Trang {issue.page_number}"
                if issue.page_number
                else issue.block_id or issue.section_title or "Chưa xác định"
            )
            document.add_heading(
                f"{index}. {issue.category.value} — {issue.severity.value}",
                level=2,
            )
            document.add_paragraph(f"Vị trí: {location}")
            document.add_paragraph(f"Nội dung gốc: {issue.original_text}")
            document.add_paragraph(f"Giải thích: {issue.explanation}")
            _add_label_value(document, "Căn cứ", issue.lesson_plan_reference)
            _add_label_value(document, "Đề xuất sửa", issue.suggested_text)
            _add_label_value(document, "Câu thay thế", issue.replacement_question)
            _add_label_value(document, "Đáp án gợi ý", issue.suggested_answer)

    document.add_heading("6. Giới hạn", level=1)
    all_limitations = list(
        dict.fromkeys(
            result.limitations
            + questions.limitations
            + answer_key.limitations
            + comparison.limitations
        )
    )
    _add_bullets(document, "", all_limitations)
    document.add_paragraph(
        "Kết quả AI chỉ mang tính hỗ trợ; giáo viên cần xác minh các mục được "
        "đánh dấu trước khi phát hành học liệu."
    )

    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()
