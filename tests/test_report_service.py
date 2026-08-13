from __future__ import annotations

import io
import json

from docx import Document

from schemas.analysis import AnalysisResult, DocumentSummary
from schemas.lesson_plan import LessonPlanProfile
from schemas.package import (
    AnswerKeyProfile,
    PackageComparisonResult,
    StudentQuestionSet,
    TeachingContextProfile,
)
from schemas.usage import AIUsageSummary
from services.report_service import (
    create_docx_report,
    create_json_report,
    create_package_docx_report,
    create_package_json_report,
)


def empty_result() -> AnalysisResult:
    return AnalysisResult(
        document_summary=DocumentSummary(
            document_type="docx",
            detected_languages=["vi"],
            summary="Tài liệu ngắn.",
        ),
        issues=[],
        general_recommendations=[],
        limitations=[],
    )


def test_json_report_is_utf8_and_valid() -> None:
    payload = json.loads(create_json_report(empty_result()).decode("utf-8"))
    assert payload["document_summary"]["summary"] == "Tài liệu ngắn."


def test_docx_report_can_be_opened() -> None:
    payload = create_docx_report(empty_result(), "lesson.docx", {})
    document = Document(io.BytesIO(payload))
    assert "BÁO CÁO KIỂM TRA HỌC LIỆU" in document.paragraphs[0].text


def test_two_document_json_report_preserves_lesson_plan_context() -> None:
    lesson_plan = LessonPlanProfile(
        source_document_type="docx",
        subject="Toán",
        topic="Phân số",
        objectives=["Nhận biết phân số."],
        summary="Giáo án về phân số.",
    )
    payload = json.loads(
        create_json_report(
            empty_result(),
            lesson_plan=lesson_plan,
            lesson_plan_file_name="giao-an.docx",
            learning_material_file_name="hoc-lieu.docx",
        ).decode("utf-8")
    )
    assert payload["lesson_plan_file_name"] == "giao-an.docx"
    assert payload["lesson_plan"]["subject"] == "Toán"
    assert "learning_material_analysis" in payload


def package_parts():
    lesson_plan = LessonPlanProfile(
        source_document_type="docx",
        subject="Tiếng Anh",
        topic="KET Reading",
        objectives=["Đọc hiểu thông tin cụ thể."],
        summary="Giáo án luyện đọc KET.",
    )
    teaching = TeachingContextProfile(
        source_file_names=["slides.pdf"],
        covered_content=["Reading Part 3"],
        key_terms=["KET"],
        worked_examples=[],
        instructions_and_conventions=["Chọn một đáp án."],
        summary="Slide hướng dẫn Reading Part 3.",
        uncertainties=[],
    )
    questions = StudentQuestionSet(
        source_file_name="practice.pdf",
        questions=[],
        summary="Bài tập thực hành.",
        limitations=[],
    )
    answer_key = AnswerKeyProfile(
        source_file_name="key.pdf",
        entries=[],
        summary="Đáp án thực hành.",
        limitations=[],
    )
    comparison = PackageComparisonResult(
        summary="Không có mục để đối chiếu.",
        comparisons=[],
        recommendations=[],
        limitations=[],
    )
    return lesson_plan, teaching, questions, answer_key, comparison


def test_package_json_preserves_all_four_roles() -> None:
    lesson_plan, teaching, questions, answer_key, comparison = package_parts()
    usage = AIUsageSummary(
        call_count=1,
        input_tokens=1_000,
        cached_input_tokens=0,
        cache_write_tokens=0,
        output_tokens=100,
        reasoning_tokens=20,
        estimated_cost_usd=0.004,
        records=[],
    )
    payload = json.loads(
        create_package_json_report(
            result=empty_result(),
            lesson_plan=lesson_plan,
            teaching_context=teaching,
            questions=questions,
            answer_key=answer_key,
            comparison=comparison,
            lesson_plan_file_name="tp.docx",
            teaching_file_names=["slides.pdf", "handout.pdf"],
            learning_material_file_name="practice.pdf",
            answer_key_file_name="key.pdf",
            usage=usage,
        ).decode("utf-8")
    )
    assert payload["package_files"] == {
        "lesson_plan": "tp.docx",
        "teaching_materials": ["slides.pdf", "handout.pdf"],
        "student_material": "practice.pdf",
        "answer_key": "key.pdf",
    }
    assert payload["teaching_context"]["key_terms"] == ["KET"]
    assert payload["ai_usage"]["estimated_cost_usd"] == 0.004


def test_package_docx_report_can_be_opened() -> None:
    lesson_plan, teaching, questions, answer_key, comparison = package_parts()
    payload = create_package_docx_report(
        result=empty_result(),
        lesson_plan=lesson_plan,
        teaching_context=teaching,
        questions=questions,
        answer_key=answer_key,
        comparison=comparison,
        lesson_plan_file_name="tp.docx",
        teaching_file_names=["slides.pdf"],
        learning_material_file_name="practice.pdf",
        answer_key_file_name="key.pdf",
        metadata={},
    )
    document = Document(io.BytesIO(payload))
    assert document.paragraphs[0].text == "BÁO CÁO KIỂM TRA BỘ HỌC LIỆU"
