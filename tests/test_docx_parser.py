from __future__ import annotations

import io

from docx import Document

from services.docx_parser import parse_docx


def make_docx(builder) -> bytes:
    document = Document()
    builder(document)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_parses_paragraph_heading_and_table_in_body_order() -> None:
    def build(document: Document) -> None:
        document.add_heading("Bài tập 1", level=1)
        document.add_paragraph("Câu 1. Chọn đáp án đúng.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Câu hỏi"
        table.cell(0, 1).text = "Đáp án"
        table.cell(1, 0).text = "2 + 2"
        table.cell(1, 1).text = "4"
        document.add_paragraph("Kết thúc")

    parsed = parse_docx(make_docx(build))
    assert [block.block_type for block in parsed.blocks] == [
        "heading",
        "paragraph",
        "table",
        "paragraph",
    ]
    assert parsed.blocks[1].section_title == "Bài tập 1"
    assert parsed.blocks[1].question_number == "1"
    assert parsed.blocks[2].rows == [["Câu hỏi", "Đáp án"], ["2 + 2", "4"]]
    assert parsed.blocks[2].block_id == "block-0003"


def test_empty_docx_has_no_blocks() -> None:
    parsed = parse_docx(make_docx(lambda _document: None))
    assert parsed.blocks == []


def test_invalid_docx_is_rejected() -> None:
    import pytest

    with pytest.raises(ValueError, match="Không thể đọc tài liệu Word"):
        parse_docx(b"broken")

